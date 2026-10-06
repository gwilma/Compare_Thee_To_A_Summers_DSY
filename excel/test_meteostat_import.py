"""End-to-end test of the DownloadMeteostatYears macro, run inside LibreOffice.

    python excel/test_meteostat_import.py [path/to/summers_dsy.xlsm]

A download folder is pre-filled with Meteostat-format files for station 03772 (2015-2021):
2018 is missing (the download attempt fails here, as LibreOffice cannot use the Windows
downloader), 2020 stops on 24 August (under 80% of the season), 2019 has three model-forecast
hours at 45 degC (dropped and interpolated) and 2017 has an 8-hour gap (left empty). The macro's
Years rows and the Years vs DSY statements are checked against the Python package.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "src"))
sys.path.insert(0, str(HERE))

from add_macros import LibreOffice, module_source  # noqa: E402
from summers_dsy.compare import years_exceeding  # noqa: E402
from summers_dsy.metrics import METRIC_INFO, AnalysisConfig, analyse  # noqa: E402
from summers_dsy.sources import meteostat as ms  # noqa: E402
from summers_dsy.synthetic import WarmSpell, synthetic_year  # noqa: E402

METRICS = list(METRIC_INFO)
STATION, LABEL = "03772", "London Heathrow Airport (03772)"
Y0, Y1 = 2016, 2021


def write_files(folder: Path) -> None:
    for year in range(2015, 2022):
        if year == 2018:
            continue
        t = synthetic_year(spells=[WarmSpell(195 + 2 * (year % 7), 3 + year % 6, 4 + 0.6 * (year % 5))],
                           warming=0.25 * (year - 2015), seed=year, year=year)
        src = pd.Series("isd_lite", index=t.index)
        if year == 2019:
            model = (t.index >= "2019-07-10 02:00") & (t.index < "2019-07-10 05:00")
            t[model], src[model] = 45.0, "metno_forecast"
        if year == 2017:
            t = t.drop(t.index[(t.index >= "2017-06-10 00:00") & (t.index < "2017-06-10 08:00")])
            src = src.reindex(t.index)
        if year == 2020:
            t, src = t[t.index < "2020-08-25"], src[src.index < "2020-08-25"]
        rows = ["year,month,day,hour,temp,temp_source,rhum,rhum_source"]
        rows += [f"{i.year},{i.month},{i.day},{i.hour},{v:.1f},{s},70,{s}" for (i, v), s in zip(t.items(), src)]
        (folder / f"{STATION}_{year}.csv").write_text("\n".join(rows) + "\n")


def run_macro(xlsm: Path, folder: Path, out: Path, out_jja: Path) -> None:
    lo = LibreOffice(port=2014)
    try:
        doc = lo.load(xlsm, macros=True)
        assert doc.BasicLibraries.getByName("VBAProject").hasByName("FolderImport")
        _, source = module_source()
        doc.BasicLibraries.getByName("Standard").insertByName("MacroTest", source)
        sheets = doc.Sheets
        m = sheets.getByName("Meteostat")
        m.getCellRangeByName("B5").setString(LABEL)
        m.getCellRangeByName("B25").setValue(Y0)
        m.getCellRangeByName("B26").setValue(Y1)
        m.getCellRangeByName("B27").setString("No")
        m.getCellRangeByName("B28").setString(str(folder))
        m.getCellRangeByName("B12").setValue(51.47)  # nearest-station helper
        m.getCellRangeByName("B13").setValue(-0.45)
        sheets.getByName("Compare").getCellRangeByName("B6").setString("Demo (synthetic)")
        doc.calculateAll()
        doc.getScriptProvider().getScript(
            "vnd.sun.star.script:Standard.MacroTest.DownloadMeteostatYearsSilent?language=Basic&location=document"
        ).invoke((), (), ())
        doc.calculateAll()
        lo.store(doc, out, "Calc Office Open XML")
        sheets.getByName("Years vs DSY").getCellRangeByName("B5").setString(METRIC_INFO["jja_mean_daily_max"][0])
        doc.calculateAll()
        lo.store(doc, out_jja, "Calc Office Open XML")
        doc.close(True)
    finally:
        lo.close()


def python_years(folder: Path, cfg: AnalysisConfig):
    """The same years analysed by the Python package from the same files."""
    def local_year(station, year):
        p = folder / f"{station}_{year}.csv"
        return ms.parse_hourly(p.read_bytes()) if p.exists() else None

    original = ms.get_year
    ms.get_year = local_year
    try:
        series = ms.fetch_years(STATION, Y0, Y1, include_model=False)
    finally:
        ms.get_year = original
    return series, {y: analyse(series.year(y), cfg) for y in range(Y0, Y1 + 1) if y not in series.meta["missing_years"]}


def check_statements(wb, year_metrics: pd.DataFrame, metric: str, failures: list) -> None:
    yv = wb["Years vs DSY"]
    lib = [r for r in wb["Library"].iter_rows(min_row=6, max_col=7 + len(METRICS), values_only=True)
           if r[0] == 1 and r[1] == "Demo (synthetic)"]
    refs = pd.DataFrame([r[7:7 + len(METRICS)] for r in lib], columns=METRICS, index=range(len(lib)))
    labels = {i: r[6] for i, r in enumerate(lib)}
    latest = yv["B10"].value
    py = years_exceeding(year_metrics, refs, labels, metric, "Demo (synthetic)", latest_complete_year=latest)
    xl = [(r[0], r[2], r[3], r[4], r[5], r[6]) for r in yv.iter_rows(min_row=15, max_row=134, max_col=7, values_only=True)
          if r[0]]
    print(f"\nYears vs DSY ({METRIC_INFO[metric][0]}): {len(xl)} rows, e.g.\n   {xl[0][5]}\n   {xl[-1][5]}")
    if len(xl) != len(py):
        failures.append(f"{metric}: {len(xl)} statement rows, Python {len(py)}")
        return
    for (lab, val, k, n, which, stmt), (_, prow) in zip(xl, py.iterrows()):
        ok = (lab == prow["label"] and np.isclose(val, prow["reference_value"]) and k == prow["years_exceeded"]
              and n == prow["years"] and (which or "") == prow["exceeding_years"] and stmt == prow["statement"])
        if not ok:
            failures.append(f"{metric}: Excel {lab!r} {k}/{n} [{which}] {stmt!r} != Python {prow['label']!r} "
                            f"{prow['years_exceeded']}/{prow['years']} [{prow['exceeding_years']}] {prow['statement']!r}")


def main(xlsm: str | None = None) -> int:
    xlsm = Path(xlsm) if xlsm else HERE / "summers_dsy.xlsm"
    failures: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp) / "meteostat"
        folder.mkdir()
        write_files(folder)
        out, out_jja = Path(tmp) / "after.xlsx", Path(tmp) / "after_jja.xlsx"
        run_macro(xlsm, folder, out, out_jja)

        wb = load_workbook(out, data_only=True, read_only=True)
        m = wb["Meteostat"]
        print("Status:", m["B31"].value)
        print("Nearest station to 51.47, -0.45:", m["B18"].value, f"{m['C18'].value:.1f} km")
        if m["B18"].value != LABEL:
            failures.append(f"nearest station {m['B18'].value!r}")
        nm = len(METRICS)
        rows = {r[0]: r for r in wb["Years"].iter_rows(min_row=6, max_row=105, max_col=8 + nm, values_only=True) if r[0]}
        thr_s, thr_t = rows[2016][6 + nm], rows[2016][7 + nm]
        cfg = AnalysisConfig(static_threshold=thr_s, twcdh_offset=thr_t)
        series, py = python_years(folder, cfg)
        print(f"\nThresholds used: SWCDH {thr_s}, TWCDH {thr_t}")
        counted = {}
        for y in range(Y0, Y1 + 1):
            r = rows.get(y)
            if r is None:
                failures.append(f"{y}: no Years row")
                continue
            cov, share, note, vals = r[3], r[4], r[5] or "", r[6:6 + nm]
            if y == 2018:
                ok = vals[0] in (None, "") and "Could not download" in note
            elif y == 2020:
                ok = vals[0] in (None, "") and np.isclose(cov, py[y].metrics["season_coverage"]) and "not counted" in note
            else:
                blank = [k for k, v in zip(METRICS, vals) if v in (None, "")]
                if blank:
                    print(y, "blank metrics in Excel:", blank, {k: py[y].metrics[k] for k in blank})
                bad = [k for k, v in zip(METRICS, vals) if v in (None, "") or not np.isclose(v, py[y].metrics[k], rtol=1e-6, atol=1e-6)]
                ok = (not bad and np.isclose(cov, py[y].metrics["season_coverage"])
                      and np.isclose(share, series.meta["model_share"][y]) and r[1] == LABEL and r[2] == STATION)
                if bad:
                    note += f" metric mismatches: {bad}"
                counted[y] = dict(zip(METRICS, vals))
            status = "OK" if ok else "FAIL"
            if not ok:
                failures.append(f"{y}: {note}")
            print(f"{status:4s} {y} coverage={cov if cov is None else round(cov, 3)} model={share} {note}")
        if 2019 in counted and counted[2019]["t_max"] >= 45:
            failures.append("2019: the model-forecast 45 degC hours were not dropped")

        year_metrics = pd.DataFrame(counted).T
        check_statements(wb, year_metrics, "swcdh", failures)
        check_statements(load_workbook(out_jja, data_only=True, read_only=True), year_metrics, "jja_mean_daily_max", failures)

    for f_ in failures:
        print("FAIL", f_)
    print("RESULT:", "PASS" if not failures else f"FAIL ({len(failures)})")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:2]))
