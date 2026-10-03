"""Check a recalculated workbook against the Python package.

    python excel/verify_workbook.py path/to/recalculated.xlsx

Reads the cached values (the workbook must have been recalculated, e.g. by opening
and saving it in Excel or LibreOffice), re-runs the same analysis in Python on the
Hourly sheet's data and thresholds, and compares the season metrics, daily comfort
temperatures and similarity ranking.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from summers_dsy.compare import similarity  # noqa: E402
from summers_dsy.metrics import DEFAULT_COMPARISON_METRICS, METRIC_INFO, AnalysisConfig, analyse  # noqa: E402
from summers_dsy.model import WeatherSeries  # noqa: E402

METRICS = list(METRIC_INFO)


def main(path: str) -> int:
    wb = load_workbook(path, data_only=True, read_only=True)
    st, hr, sm, lb, cp, dy = (wb[n] for n in ("Settings", "Hourly", "Summary", "Library", "Compare", "Daily"))

    s = {r[0].value: r[1].value for r in st.iter_rows(min_row=4, max_row=60, max_col=2) if r[0].value}
    swcdh, twcdh = s["SWCDH threshold used (°C)"], s["TWCDH offset used (K)"]
    typical = s["Typical/design year file? (wrap running mean)"] == "Yes"

    rows = [r for r in hr.iter_rows(min_row=2, max_col=5, values_only=True) if r[4] is not None]
    df = pd.DataFrame(rows, columns=["year", "month", "day", "hour", "t"])
    year = df["year"].fillna(s["Nominal year for rows with no Year"]).astype(int)
    idx = pd.to_datetime({"year": year, "month": df["month"], "day": df["day"], "hour": df["hour"] - 1})
    series = WeatherSeries("wb", pd.DataFrame({"dry_bulb": df["t"].to_numpy()}, index=idx), meta={"typical_year": typical})
    cfg = AnalysisConfig(static_threshold=swcdh, twcdh_offset=twcdh,
                         hot_day_threshold=s["Hot day: daily max ≥ (°C)"], warm_night_threshold=s["Warm night: daily min ≥ (°C)"],
                         night_start_hour=int(s["Night starts (hour, 0–23)"]), night_end_hour=int(s["Night ends (hour, 0–23)"]),
                         night_threshold=s["Night mean threshold (°C)"],
                         night_min_hours=int(s["Minimum valid hours for a night mean"]))
    py = analyse(series, cfg)

    failures = 0
    print(f"{'metric':32s} {'excel':>14s} {'python':>14s}")
    xl_metrics = {}
    for i, key in enumerate(METRICS):
        v = sm.cell(row=5 + i, column=2).value
        xl_metrics[key] = v
        ok = np.isclose(float(v), py.metrics[key], rtol=1e-6, atol=1e-6)
        failures += not ok
        print(f"{key:32s} {float(v):14.4f} {py.metrics[key]:14.4f} {'' if ok else '  <-- MISMATCH'}")

    # Daily comfort temperature over the season
    drows = [r for r in dy.iter_rows(min_row=2, max_col=12, values_only=True) if r[1] is not None and r[1] != ""]
    xl_tc = pd.Series({pd.Timestamp(r[1]).normalize(): r[9] for r in drows if r[11] == 1 and r[9] not in (None, "")})
    py_tc = py.season_daily()["t_comf"]
    diff = (xl_tc - py_tc.reindex(xl_tc.index)).abs().max()
    print(f"\nmax |Tcomf excel - python| over season days: {diff:.2e}")
    failures += not diff < 1e-6

    # Similarity ranking
    lib = [r for r in lb.iter_rows(min_row=6, max_col=7 + len(METRICS), values_only=True) if r[0] == 1]
    ref = pd.DataFrame([r[7:] for r in lib], columns=METRICS, index=[r[6] for r in lib]).astype(float)
    weights = {k: cp.cell(row=13 + j, column=8).value for j, k in enumerate(METRICS)}
    used = [k for k in METRICS if weights[k]]
    ranked = similarity(xl_metrics, ref, used, weights)
    xl_top = [(cp.cell(row=12 + k, column=2).value, cp.cell(row=12 + k, column=4).value) for k in range(1, 11)]
    print("\nrank  excel label / score                                   python label / score")
    for k, ((xl_lab, xl_score), (py_lab, py_row)) in enumerate(zip(xl_top, ranked.head(10).iterrows()), start=1):
        ok = xl_lab == py_lab and np.isclose(xl_score, py_row["score"], atol=1e-6)
        failures += not ok
        print(f"{k:>4}  {xl_lab:45s} {xl_score:7.3f}   {py_lab:45s} {py_row['score']:7.3f}{'' if ok else '  <-- MISMATCH'}")
    print("\nHeadline:", cp["B4"].value)
    print("WCDH context:", cp["B5"].value)
    print("\nDefault metrics used:", used == DEFAULT_COMPARISON_METRICS)
    print("RESULT:", "PASS" if failures == 0 else f"FAIL ({failures} mismatches)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
