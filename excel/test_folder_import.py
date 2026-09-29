"""End-to-end test of the ImportWeatherFolder macro, run inside LibreOffice.

    python excel/test_folder_import.py [path/to/summers_dsy.xlsm]

Builds a folder tree of synthetic weather files (CIBSE 2025 and 2016 naming, EPW and CSV,
plus files that should be skipped), runs the macro through Python-UNO, saves the result and
checks every imported Library row against the Python package: file metadata, the per-location
thresholds and all 18 metrics.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "src"))
sys.path.insert(0, str(HERE))

from add_macros import LibreOffice, module_source  # noqa: E402
from summers_dsy.io import load_weather_file  # noqa: E402
from summers_dsy.metrics import METRIC_INFO, AnalysisConfig, analyse, static_threshold, threshold_offset  # noqa: E402
from summers_dsy.synthetic import WarmSpell, synthetic_series, to_epw  # noqa: E402

METRICS = list(METRIC_INFO)

# relative path -> (spells, warming, seed, expected (location, kind, period, emissions, percentile))
FILES = {
    "CIBSE 2025/Z1/Z1_TRY_2030s_HIGH50_CIBSE_v1.1.epw": ([WarmSpell(200, 4, 3)], 0.0, 1, ("Zone 1", "TRY", "2030s", "High", 50)),
    "CIBSE 2025/Z1/Z1_DSY1_2030s_HIGH50_CIBSE_v1.1.epw": ([WarmSpell(195, 6, 5)], 0.3, 2, ("Zone 1", "DSY1", "2030s", "High", 50)),
    "CIBSE 2025/Z1/Z1_DSY2_2050s_HIGH10_CIBSE_v1.1.epw": ([WarmSpell(215, 5, 9)], 1.5, 3, ("Zone 1", "DSY2", "2050s", "High", 10)),
    "CIBSE 2025/Z1/2080s/Z1_DSY3_2080s_LOW90_CIBSE_v1.1.epw": ([WarmSpell(205, 20, 6)], 2.0, 4, ("Zone 1", "DSY3", "2080s", "Low", 90)),
    "CIBSE 2025/Z7/Z7_TRY_2030s_HIGH50_CIBSE_v1.1.epw": ([WarmSpell(200, 4, 2)], -1.0, 5, ("Zone 7", "TRY", "2030s", "High", 50)),
    "CIBSE 2025/Z7/Z7_DSY1_2050s_MEDIUM10_CIBSE_v1.1.epw": ([WarmSpell(210, 7, 6)], 0.5, 6, ("Zone 7", "DSY1", "2050s", "Medium", 10)),
    "CIBSE 2016/London/London_LHR_TRY.epw": ([WarmSpell(205, 3, 3)], 0.4, 7, ("London (Heathrow)", "TRY", "Baseline", None, None)),
    "CIBSE 2016/London/London_LHR_DSY1_2050High50.epw": ([WarmSpell(200, 8, 6)], 1.4, 8, ("London (Heathrow)", "DSY1", "2050s", "High", 50)),
    "CIBSE 2016/Manchester/Manchester_DSY2_2080s_High_90.csv": ([WarmSpell(220, 4, 10)], 2.5, 9, ("Manchester", "DSY2", "2080s", "High", 90)),
}
SKIPPED = ["junk/no_header.csv", "junk/notes.txt"]


def make_tree(root: Path) -> None:
    for rel, (spells, warming, seed, _) in FILES.items():
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        series = synthetic_series(rel, spells=spells, warming=warming, seed=seed)
        if rel.endswith(".epw"):
            path.write_text(to_epw(series, year=1990 + seed).replace("\n", "\r\n"))  # CIBSE files use CRLF
        else:
            rows = ["CIBSE-style test file", "Month,Day,Hour,Dry bulb temperature (C),RH (%)"]
            rows += [f"{ts.month},{ts.day},{ts.hour + 1},{t:.1f},70" for ts, t in series.dry_bulb.items()]
            path.write_text("\n".join(rows) + "\n")
    (root / "junk").mkdir(exist_ok=True)
    (root / "junk/no_header.csv").write_text("1,2,3\n4,5,6\n")
    (root / "junk/notes.txt").write_text("not a weather file\n")
    # An observed year with real dates: imported as file type "Other"; "heathrow" in its name sets the location.
    obs = synthetic_series("obs", typical=False, year=2019, spells=[WarmSpell(206, 5, 8)], warming=0.8, seed=11)
    (root / "observed").mkdir(exist_ok=True)
    (root / "observed/heathrow_2019.epw").write_text(to_epw(obs))


def run_macro(xlsm: Path, folder: Path, out: Path) -> None:
    lo = LibreOffice(port=2003)
    try:
        doc = lo.load(xlsm, macros=True)
        # Check the workbook really carries the module, then run the same source from the Standard
        # library: LibreOffice will not execute an imported VBA project's modules over UNO.
        vba = doc.BasicLibraries.getByName("VBAProject")
        assert vba.hasByName("FolderImport"), "the .xlsm has no FolderImport module"
        _, source = module_source()
        doc.BasicLibraries.getByName("Standard").insertByName("FolderImportTest", source)
        script = doc.getScriptProvider().getScript(
            "vnd.sun.star.script:Standard.FolderImportTest.ImportWeatherFolderFrom?language=Basic&location=document")
        script.invoke((str(folder),), (), ())
        script.invoke((str(folder),), (), ())  # a second run must replace rows, not add duplicates
        # Before: Settings hold the demo thresholds, so comparing with Zone 1 should raise the warning.
        doc.calculateAll()
        compare = doc.Sheets.getByName("Compare")
        compare.getCellRangeByName("B6").setString("Zone 1")
        doc.calculateAll()
        warning_before = compare.getCellRangeByName("B7").getString()
        doc.getScriptProvider().getScript(
            "vnd.sun.star.script:Standard.FolderImportTest.UseLibraryThresholdsSilent?language=Basic&location=document"
        ).invoke((), (), ())
        doc.calculateAll()
        warning_after = compare.getCellRangeByName("B7").getString()
        headline = compare.getCellRangeByName("B4").getString()
        (out.parent / "warnings.txt").write_text(f"{warning_before}\n{warning_after}\n{headline}\n")
        lo.store(doc, out, "Calc Office Open XML")
        doc.close(True)
    finally:
        lo.close()


def main(xlsm: str | None = None) -> int:
    xlsm = Path(xlsm) if xlsm else HERE / "summers_dsy.xlsm"
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp) / "weather"
        make_tree(root)
        out = Path(tmp) / "after.xlsx"
        run_macro(xlsm, root, out)

        wb = load_workbook(out, data_only=True, read_only=True)
        lib = [r for r in wb["Library"].iter_rows(min_row=6, max_col=28, values_only=True) if r[27]]
        log = [r for r in wb["Import log"].iter_rows(min_row=4, max_col=9, values_only=True) if r[1]]
        print("Import log:")
        for r in log:
            print("  ", (Path(r[0]).name if r[0] else ""), "|", r[1], "|", r[2], "|", r[8] or "")

        by_file = {Path(r[27]).name: r for r in lib if r[27] != "synthetic demo data"}
        expected_files = {Path(rel).name: v for rel, v in FILES.items()}
        expected_files["heathrow_2019.epw"] = (None, None, None, 0, ("London (Heathrow)", "Other", "Baseline", None, None))

        # Thresholds per location, from each location's TRY
        thr = {}
        for rel, (_, _, _, meta) in FILES.items():
            if meta[1] == "TRY":
                s = load_weather_file((root / rel).read_bytes(), Path(rel).name)
                cfg = AnalysisConfig()
                thr[meta[0]] = (round(static_threshold(s, cfg), 3), round(threshold_offset(s, cfg), 3))

        for fname, spec in expected_files.items():
            meta = spec[-1]
            row = by_file.get(fname)
            if row is None:
                print(f"MISSING  {fname}")
                failures += 1
                continue
            got = (row[1], row[2], row[3], row[4], row[5])
            ok_meta = got == meta
            path = next(root.rglob(fname))
            series = load_weather_file(path.read_bytes(), fname)
            if meta[1] == "Other":
                series.meta["typical_year"] = False
            loc_thr = thr.get(meta[0])
            if loc_thr:
                ok_thr = np.isclose(row[25], loc_thr[0]) and np.isclose(row[26], loc_thr[1])
            else:
                ok_thr = True  # no TRY for this location: Settings values are used
            py = analyse(series, AnalysisConfig(static_threshold=row[25], twcdh_offset=row[26])).metrics
            bad = [k for j, k in enumerate(METRICS) if not np.isclose(row[7 + j], py[k], rtol=1e-6, atol=1e-6)]
            status = "OK" if ok_meta and ok_thr and not bad else "FAIL"
            failures += status == "FAIL"
            print(f"{status:4s} {fname:44s} meta={got} thr=({row[25]:.3f},{row[26]:.3f})"
                  + ("" if ok_meta else f" expected {meta}") + ("" if ok_thr else f" expected thr {loc_thr}")
                  + (f" metric mismatches: {bad}" if bad else ""))

        skipped = {Path(r[0]).name for r in log if r[1] == "Skipped" and r[0]}
        if "no_header.csv" not in skipped:
            print("FAIL no_header.csv was not reported as skipped")
            failures += 1
        if any("notes.txt" in (r[0] or "") for r in log):
            print("FAIL notes.txt should have been ignored")
            failures += 1

        before, after, headline = (out.parent / "warnings.txt").read_text().split("\n")[:3]
        print("Warning before UseLibraryThresholds:", before[:90] + "…" if before else "(none)")
        print("Warning after:", after or "(none)")
        print("Headline for Zone 1:", headline)
        if not before or after:
            print("FAIL threshold warning did not behave as expected")
            failures += 1
        imported = [r for r in lib if r[27] != "synthetic demo data"]
        if len(imported) != 10:
            print(f"FAIL expected 10 imported rows after two runs, found {len(imported)}")
            failures += 1

        # The user's own data and settings are put back after the import.
        st = {r[0]: r[1] for r in wb["Settings"].iter_rows(min_row=4, max_row=45, max_col=2, values_only=True) if r[0]}
        hr_first = next(wb["Hourly"].iter_rows(min_row=2, max_row=2, max_col=5, values_only=True))
        restored = hr_first[0] == 2020 and st["Typical/design year file? (wrap running mean)"] == "No"
        z1 = thr["Zone 1"]
        print("Settings thresholds after UseLibraryThresholds:", st["SWCDH threshold override (°C)"], st["TWCDH offset override (K)"], "expected", z1)
        if not (np.isclose(st["SWCDH threshold override (°C)"], z1[0]) and np.isclose(st["TWCDH offset override (K)"], z1[1])):
            failures += 1
        print("Hourly data and Settings restored:", restored)
        failures += not restored
        demo_rows = sum(1 for r in lib if r[27] == "synthetic demo data")
        print("Demo rows kept:", demo_rows)

    print("RESULT:", "PASS" if failures == 0 else f"FAIL ({failures})")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:2]))
